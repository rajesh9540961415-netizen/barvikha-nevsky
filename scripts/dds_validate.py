#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Сверка разобранных данных с собственными итогами Excel-файла.

Для каждого месяца берём из листа его же строки-итоги («Расходы ИТОГО»,
заголовки блоков «Списания с р/с» и т.д.) и сравниваем с суммой,
которую посчитал парсер. Расхождение = либо ошибка парсера, либо
битая формула в исходном файле.
"""
import csv
import datetime as dt
import openpyxl
from pathlib import Path
from collections import defaultdict
from dds_import import SHEETS, SRC, OUT, find_header, norm, is_num, SOURCES, STOP

wb = openpyxl.load_workbook(SRC, data_only=True, read_only=True)

tx = defaultdict(float)          # (sheet, source) -> сумма по парсеру
tx_month = defaultdict(float)    # sheet -> сумма по парсеру
with open(f"{OUT}/transactions.csv", encoding="utf-8") as f:
    for r in csv.DictReader(f):
        tx[(r["sheet"], r["source"])] += float(r["amount"])
        tx_month[r["sheet"]] += float(r["amount"])

met = defaultdict(float)
with open(f"{OUT}/daily_metrics.csv", encoding="utf-8") as f:
    for r in csv.DictReader(f):
        met[(r["sheet"], r["metric"])] += float(r["value"])

print(f"{'лист':<12} {'расход (парсер)':>16} {'расход (лист)':>16} {'Δ':>12}   "
      f"{'выручка (парсер)':>17} {'выручка (лист)':>16} {'Δ':>10}  ошибки формул")
print("─" * 130)

errs_total = 0
for name, y, m in SHEETS:
    ws = wb[name]
    rows = [list(r) for r in ws.iter_rows(min_row=1, max_row=ws.max_row,
                                          max_col=min(ws.max_column, 60), values_only=True)]
    hrow, lcol, dates = find_header(ws, rows)
    day_cols = set(dates)

    sheet_exp = None      # итог расходов по строке листа
    sheet_rev = None
    formula_errs = 0
    block_tot = {}

    for i, row in enumerate(rows, start=1):
        for v in row:
            if isinstance(v, str) and v.startswith("#") and v.endswith(("!", "?", "0!", "A")):
                formula_errs += 1
        lab = norm(row[lcol - 1]) if len(row) >= lcol else ""
        vals = [row[c - 1] for c in sorted(day_cols) if len(row) >= c and is_num(row[c - 1])]
        s = sum(vals)
        if lab.startswith("расходы  итого") or lab.startswith("расходы итого"):
            sheet_exp = s
        if lab == "выручка итого":
            sheet_rev = s
        matched = next((src for pref, src in SOURCES if lab.startswith(pref)), None)
        if matched:
            block_tot[matched] = block_tot.get(matched, 0) + s

    errs_total += formula_errs
    p_exp = tx_month.get(name, 0)
    p_rev = met.get((name, "revenue_total"), 0)
    d_exp = (p_exp - sheet_exp) if sheet_exp is not None else None
    d_rev = (p_rev - sheet_rev) if sheet_rev is not None else None

    f = lambda v: "—" if v is None else f"{v:>16,.0f}".replace(",", " ")
    g = lambda v: "—" if v is None else ("ok" if abs(v) < 1 else f"{v:,.0f}".replace(",", " "))
    print(f"{name:<12} {p_exp:>16,.0f} {f(sheet_exp)} {g(d_exp):>12}   "
          f"{p_rev:>17,.0f} {f(sheet_rev)} {g(d_rev):>10}  {formula_errs or ''}"
          .replace(",", " "))

    for src, tot in sorted(block_tot.items()):
        mine = tx.get((name, src), 0)
        if abs(mine - tot) >= 1:
            print(f"{'':<12}   ↳ блок «{src}»: парсер {mine:,.0f} / лист {tot:,.0f} "
                  f"→ Δ {mine - tot:,.0f}".replace(",", " "))

print("─" * 130)
print(f"всего ячеек с ошибками формул в помесячных листах: {errs_total}")
