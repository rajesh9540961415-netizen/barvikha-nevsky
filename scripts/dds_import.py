#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Разбор ручного файла ДДС в плоские таблицы.

Читает помесячные листы Excel и раскладывает их в два CSV:
  daily_metrics.csv  — дневные показатели (выручка, чеки, столы)
  transactions.csv   — каждая оплата: дата, юрлицо, источник денег, контрагент, статья, сумма

Парсер ориентируется на подписи строк, а не на их номера: за два года
структура листа менялась трижды, номера строк съезжали.
"""
import csv
import datetime as dt
import openpyxl
from pathlib import Path
from collections import defaultdict

# Актуальный файл — «ДДС апрель мой 2026»: в нём пересмотрены январь–март 2026
# (проставлены статьи, добавлены наличные) и есть апрель. Прежний файл оставлен
# как запасной для месяцев, которых в новом нет.
SRC = "/Users/user/Downloads/Финансы/ДДС апрель мой 2026.xlsx"
SRC_FALLBACK = "/Users/user/Downloads/ДДС  2026.xlsx"
OUT = str(Path(__file__).resolve().parent.parent / "data" / "dds")

# лист -> (год, месяц); порядок = хронологический
SHEETS = [
    ("Май", 2024, 5), ("Июнь", 2024, 6), ("Июль", 2024, 7), ("Август", 2024, 8),
    ("Сентябрь", 2024, 9), ("Октябрь", 2024, 10), ("Ноябрь", 2024, 11), ("Декабрь", 2024, 12),
    ("ЯНВАРЬ25", 2025, 1), ("ФЕВРАЛЬ25", 2025, 2), ("МАРТ25", 2025, 3), ("АПРЕЛЬ25", 2025, 4),
    ("МАЙ25", 2025, 5), ("ИЮНЬ25", 2025, 6), ("ИЮЛЬ25", 2025, 7), ("АВГУСТ25", 2025, 8),
    ("СЕНТЯБРЬ25", 2025, 9), ("ОКТЯБРЬ25", 2025, 10), ("НОЯБРЬ25", 2025, 11), ("ДЕКАБРЬ25", 2025, 12),
    ("ЯНВАРЬ26", 2026, 1), ("ФЕВРАЛЬ26", 2026, 2), ("МАРТ26", 2026, 3), ("АПРЕЛЬ26", 2026, 4),
]

# подпись строки -> каноническое имя показателя
METRICS = {
    "выручка ип борзова в.в.": "revenue_cashless",
    "поступление ип борзова в.в.": "revenue_cashless",
    "поступление ип хисамов т.м.": "revenue_cashless",
    'поступления ооо "мята"': "revenue_cashless_myata",
    'выручка ооо "мята"': "revenue_cashless_myata",
    'поступления ооо "мав"': "revenue_cashless",
    "итого безнал": "revenue_cashless_total",
    "наличка ип борзова в.в.": "revenue_cash",
    "наличка ип хисамов т.м.": "revenue_cash",
    'наличка ооо "мята"': "revenue_cash_myata",
    'наличка ооо "мав"': "revenue_cash",
    "итого наличка": "revenue_cash_total",
    "выручка итого": "revenue_total",
    "выручка прогноз": "revenue_forecast",
    "количество столов": "tables",
    "средний чек по стола": "avg_check_table",
    "количество чеков": "checks",
    "средний чек по чекам": "avg_check_receipt",
}

# подпись строки -> источник денег
SOURCES = [
    ("расходы наличкой", "cash"),
    ("списание с бизнес карты", "card"),
    ("списания с бизнес карты", "card"),
    ("списания с р/с", "bank"),
    ("списание с р/с", "bank"),
]

# начало блока со сводкой по статьям
STRUCT_HEADERS = ("расходы ооо", "структура расходов", "расходы  итого", "расходы итого")
STOP = ("проверка", "прверка")


def norm(v):
    return str(v).strip().lower().replace("\xa0", " ") if v is not None else ""


def is_num(v):
    return isinstance(v, (int, float)) and not isinstance(v, bool)


def find_header(ws, rows):
    """Возвращает (row_idx, label_col, {col: date}). Ищем строку с датами."""
    best = None
    for i, row in enumerate(rows[:12], start=1):
        dates = {}
        for c, v in enumerate(row, start=1):
            if isinstance(v, dt.datetime):
                dates[c] = v.date()
        if len(dates) >= 25 and (best is None or len(dates) > len(best[2])):
            first = min(dates)
            best = (i, first - 1, dates)
    return best


def parse_sheet(wb, name, year, month, tx, met, issues):
    ws = wb[name]
    rows = [list(r) for r in ws.iter_rows(min_row=1, max_row=ws.max_row,
                                          max_col=min(ws.max_column, 60), values_only=True)]
    hdr = find_header(ws, rows)
    if not hdr:
        issues.append((name, "-", "нет строки с датами", ""))
        return
    hrow, lcol, dates = hdr

    # даты должны попадать в свой месяц — иначе лист собран из чужого шаблона
    bad = [str(d) for d in dates.values() if (d.year, d.month) != (year, month)]
    if bad:
        issues.append((name, f"строка {hrow}", "даты не из своего месяца", ", ".join(bad[:5])))

    def label(row):
        return norm(row[lcol - 1]) if len(row) >= lcol else ""

    # ── дневные показатели ────────────────────────────────────────────────
    # Блок выручки — это строки от шапки до «Средний чек по чекам». Ниже идут
    # банковские поступления, там подписи повторяются и в выручку не годятся.
    last = next((i for i, row in enumerate(rows, start=1)
                 if i > hrow and label(row) == "средний чек по чекам"), len(rows))
    for i, row in enumerate(rows[:last], start=1):
        if i <= hrow:
            continue
        metric = METRICS.get(label(row))
        if not metric:
            continue
        for c, d in dates.items():
            v = row[c - 1] if len(row) >= c else None
            if is_num(v):
                met.append({"date": d, "sheet": name, "metric": metric,
                            "value": round(float(v), 2)})

    # ── платежи по контрагентам ───────────────────────────────────────────
    source = None
    entity_hint = None
    for i, row in enumerate(rows, start=1):
        if i <= hrow:
            continue
        lab = label(row)
        cat_cell = row[0] if row else None       # столбец A — статья расхода
        ent_cell = row[1] if len(row) > 1 else None

        if lab.startswith(STOP):
            source = None
            continue
        matched = next((s for pref, s in SOURCES if lab.startswith(pref)), None)
        if matched:
            source = matched
            entity_hint = None
            if "мав" in lab:
                entity_hint = "ООО МАВ"
            elif "ип" in lab:
                entity_hint = "ИП"
            continue
        if source is None:
            continue
        if cat_cell is None:
            # строка контрагента, у которой забыли заполнить статью расхода
            has_money = any(is_num(row[c - 1]) and abs(row[c - 1]) > 0.004
                            for c in dates if len(row) >= c)
            if not (has_money and lab and not lab.startswith(STRUCT_HEADERS)):
                continue
            issues.append((name, f"строка {i}", "нет статьи расхода", str(row[lcol - 1])[:40]))

        category = str(cat_cell).strip() if cat_cell is not None else "БЕЗ СТАТЬИ"
        name_cp = str(row[lcol - 1]).strip() if len(row) >= lcol and row[lcol - 1] is not None else ""
        if not name_cp:
            continue
        entity = None
        if isinstance(ent_cell, str) and ent_cell.strip() and not ent_cell.strip().isdigit():
            entity = ent_cell.strip()
        if entity is None:
            entity = entity_hint or ("наличные" if source == "cash" else "")

        for c, d in dates.items():
            v = row[c - 1] if len(row) >= c else None
            if is_num(v) and abs(v) > 0.004:
                tx.append({"date": d, "sheet": name, "entity": entity, "source": source,
                           "counterparty": name_cp, "category": category,
                           "amount": round(float(v), 2)})


def main():
    wb = openpyxl.load_workbook(SRC, data_only=True, read_only=True)
    wb2 = None
    tx, met, issues = [], [], []
    for name, y, m in SHEETS:
        if name in wb.sheetnames:
            parse_sheet(wb, name, y, m, tx, met, issues)
            continue
        if wb2 is None:
            wb2 = openpyxl.load_workbook(SRC_FALLBACK, data_only=True, read_only=True)
        if name in wb2.sheetnames:
            issues.append((name, "-", "лист взят из запасного файла", ""))
            parse_sheet(wb2, name, y, m, tx, met, issues)
        else:
            issues.append((name, "-", "листа нет ни в одном файле", ""))

    tx.sort(key=lambda r: (r["date"], r["source"], r["counterparty"]))
    met.sort(key=lambda r: (r["date"], r["metric"]))

    with open(f"{OUT}/transactions.csv", "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, ["date", "sheet", "entity", "source", "counterparty", "category", "amount"])
        w.writeheader()
        w.writerows(tx)
    with open(f"{OUT}/daily_metrics.csv", "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, ["date", "sheet", "metric", "value"])
        w.writeheader()
        w.writerows(met)

    print(f"платежей: {len(tx)}   строк показателей: {len(met)}")
    print(f"период платежей: {tx[0]['date']} … {tx[-1]['date']}")
    print(f"сумма расходов всего: {sum(r['amount'] for r in tx):,.0f} ₽".replace(",", " "))
    print()
    by_src = defaultdict(float)
    for r in tx:
        by_src[r["source"]] += r["amount"]
    print("по источникам:")
    for k, v in sorted(by_src.items(), key=lambda x: -x[1]):
        print(f"   {k:6} {v:>16,.0f}".replace(",", " "))
    print()
    by_year = defaultdict(float)
    for r in tx:
        by_year[r["date"].year] += r["amount"]
    print("по годам:", {k: f"{v:,.0f}".replace(",", " ") for k, v in sorted(by_year.items())})
    print(f"\nуникальных контрагентов: {len({r['counterparty'] for r in tx})}")
    print(f"уникальных статей: {len({r['category'] for r in tx})}")
    if issues:
        print("\nЗАМЕЧАНИЯ ПРИ РАЗБОРЕ:")
        for a, b, c, d in issues:
            print(f"   {a} | {b} | {c} | {d}")


if __name__ == "__main__":
    main()
