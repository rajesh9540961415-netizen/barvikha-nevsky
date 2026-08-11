#!/usr/bin/env python3
"""Вписывает выручку в отчёт по чаевым — из обеих касс Quick Resto.

Выручка нужна отчёту ради строки «% чаевых от выручки». Без неё та показывает
0,0 %, и весь смысл сравнения пропадает: чаевые сами по себе ни о чём не
говорят, важна их доля от того, что заведение заработало.

⚠️ **Касс две, и складывать надо обе.** С 7 июля 2026 в заведении работают
два юрлица со своими облаками Quick Resto: ООО «МАВ» и ИП Шалаева. Выручка
заведения — это сумма. Взять одну кассу значит занизить выручку почти вдвое
и получить вдвое завышенный процент чаевых.

День берётся по дате ОТКРЫТИЯ смены — так же, как в приложении смены и в
ручной таблице. Бизнес-день в Quick Resto начинается в 11:00, приложение
настроено так же, поэтому дни в обоих источниках означают одно и то же.

Запуск:
    python3 scripts/report_revenue.py              # месяц вчерашнего дня
    python3 scripts/report_revenue.py 2026-08      # конкретный месяц

Доступы берёт из `.env`: QR_MAV_* и QR_SHALAEVA_*.
"""

from __future__ import annotations

import sys
from collections import defaultdict
from datetime import date, timedelta
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

from openpyxl import load_workbook  # noqa: E402

from qr_sales import CLOUDS, business_day, fetch_shifts, read_env  # noqa: E402
from report_path import describe, report_dir  # noqa: E402

ROOT = Path(__file__).resolve().parent.parent
МЕСЯЦЫ = ("январь февраль март апрель май июнь июль "
          "август сентябрь октябрь ноябрь декабрь").split()


def выручка_по_дням(месяц: str) -> tuple[dict[date, float], dict[str, float]]:
    """Сумма по обеим кассам за месяц. Возвращает (по дням, итоги по кассам)."""
    по_дням: dict[date, float] = defaultdict(float)
    по_кассам: dict[str, float] = {}

    for cloud, (_, название) in CLOUDS.items():
        cfg = read_env(cloud)
        всего = 0.0
        for s in fetch_shifts(cfg):
            d = business_day(s)
            # Открытые смены не берём: суммы в них ещё меняются, и вчерашний
            # день попал бы в отчёт недосчитанным.
            if not d or s.get("status") != "CLOSED":
                continue
            if d.strftime("%Y-%m") != месяц:
                continue
            нал = (s.get("totalCash") or 0) - (s.get("totalReturnCash") or 0)
            безнал = (s.get("totalCard") or 0) - (s.get("totalReturnCard") or 0)
            по_дням[d] += нал + безнал
            всего += нал + безнал
        по_кассам[название] = всего
        print(f"  {название:<14} {всего:>14,.0f} ₽".replace(",", " "))

    return по_дням, по_кассам


def main() -> None:
    месяц = sys.argv[1] if len(sys.argv) > 1 else (
        date.today() - timedelta(days=1)).strftime("%Y-%m")
    год, мес = (int(x) for x in месяц.split("-"))

    имя = f"Чаевые-дубль-{МЕСЯЦЫ[мес - 1]}-{год}.xlsx"
    найдено = [p for p in (report_dir() / имя, ROOT / имя) if p.exists()]
    if not найдено:
        sys.exit(
            f"Не нашёл отчёт «{имя}».\n"
            f"  Создать: python3 scripts/report_new.py {месяц}"
        )
    отчёт = найдено[0]

    print(f"Выручка за {МЕСЯЦЫ[мес - 1]} {год}:")
    по_дням, по_кассам = выручка_по_дням(месяц)
    if not по_дням:
        sys.exit("✗ Ни одной закрытой смены за этот месяц. В отчёт писать нечего.")

    wb = load_workbook(отчёт)
    ws = wb[wb.sheetnames[0]]

    строка = next(
        (r for r in range(1, 40) if ws.cell(r, 1).value == "Выручка"), None
    )
    if строка is None:
        sys.exit("В отчёте нет строки «Выручка». Отчёт собран старым шаблоном.")

    записано = 0
    for д, сумма in sorted(по_дням.items()):
        нужно = д.strftime("%d.%m.%Y")
        столбец = next(
            (c for c in range(2, ws.max_column + 1)
             if str(ws.cell(1, c).value or "").strip() == нужно),
            None,
        )
        if столбец is None:
            print(f"  ⚠ в отчёте нет столбца {нужно} — пропускаю")
            continue
        ws.cell(строка, столбец, round(сумма, 2))
        записано += 1

    wb.save(отчёт)

    итого = sum(по_кассам.values())
    print(f"  {'ВСЕГО':<14} {итого:>14,.0f} ₽".replace(",", " "))
    print(f"\nЗаписано дней: {записано}")
    print(f"Отчёт: {отчёт.name}")
    print(f"Папка: {describe(отчёт.parent)}")
    print("\nПроцент чаевых от выручки таблица посчитает сама при открытии.")


if __name__ == "__main__":
    main()
