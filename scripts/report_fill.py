#!/usr/bin/env python3
"""Вписывает выгрузку смены из приложения в отчёт.

Приложение отдаёт файл вида «смена-2026-08-01.json» (кнопка «Выгрузить
смену для отчёта» на экране итогов). Скрипт находит в отчёте столбец
нужной даты и раскладывает суммы по блокам.

Запуск:
    python3 scripts/report_fill.py смена-2026-08-01.json
    python3 scripts/report_fill.py смена-2026-08-01.json --report Чаевые-дубль-август-2026.xlsx

Существующие значения в этом столбце перезаписываются: смену можно
выгружать повторно, если данные поправили.
"""

import argparse
import json
import sys
from pathlib import Path

from openpyxl import load_workbook

from report_path import describe, report_dir

ROOT = Path(__file__).resolve().parent.parent

# Блок отчёта → что из выгрузки в него кладём.
#
# Поля `staff` (удержано за питание) здесь нет намеренно: отчёт фиксирует
# выплаты В ПОЛЬЗУ сотрудника, а обед — его встречная покупка у заведения.
# Решение владельца 08.08.2026.
WAITER_BLOCKS = {
    "ОФИЦИАНТЫ · ЧАЕВЫЕ НАЛИЧНЫЕ": "cash",
    "ОФИЦИАНТЫ · ЧАЕВЫЕ ТЕРМИНАЛ": "terminal",
    "ОФИЦИАНТЫ · ЧАЕВЫЕ QR": "qr",
    "ОФИЦИАНТЫ · СЕРВИСНЫЙ СБОР": "service",
    # Блок зовётся «начислено», а не «на руки»: питание из этой цифры
    # не вычтено, и поевший получит на руки меньше.
    "ОФИЦИАНТЫ · НАЧИСЛЕНО": "accrued",
}


def value_of(waiter: dict, field: str) -> float:
    """Сумма из выгрузки. Выручает, когда выгрузка старше формата.

    Выгрузки до 08.08.2026 звали начисленное полем `net`, и питание из него
    было вычтено. Берём как есть — другого числа в файле просто нет, а
    остановиться значило бы не заполнить смену вовсе.
    """
    if field == "accrued" and "accrued" not in waiter:
        return waiter.get("net", 0)
    return waiter.get(field, 0)
ADMIN_BLOCK = "АДМИНИСТРАТОРЫ · на руки (доля 20 %)"
CASHIER_BLOCK = "КАССИРЫ · на руки (доля от 750 ₽)"


def is_header(value) -> bool:
    """Заголовок блока: первое слово заглавными («АДМИНИСТРАТОРЫ · на руки»)."""
    if not isinstance(value, str) or not value.strip():
        return False
    first = value.split()[0].strip("·")
    return first.isupper() and len(first) > 2


def find_blocks(ws) -> dict[str, tuple[int, int]]:
    """Границы блоков: от строки после заголовка до следующего заголовка."""
    heads = [
        (r, ws.cell(r, 1).value)
        for r in range(1, ws.max_row + 1)
        if is_header(ws.cell(r, 1).value)
        and ws.cell(r, 1).value != "СВОДНАЯ ТАБЛИЦА"
    ]
    blocks = {}
    for i, (row, title) in enumerate(heads):
        end = heads[i + 1][0] - 2 if i + 1 < len(heads) else ws.max_row
        blocks[title] = (row + 1, end)
    return blocks


def find_date_column(ws, iso_day: str) -> int:
    """Столбец с датой смены. В отчёте даты записаны как 01.08.2026."""
    year, month, day = iso_day.split("-")
    wanted = f"{day}.{month}.{year}"
    for col in range(2, ws.max_column + 1):
        if str(ws.cell(1, col).value or "").strip() == wanted:
            return col
    sys.exit(f"В отчёте нет столбца с датой {wanted}. Тот ли это месяц?")


def put(ws, block: tuple[int, int], col: int, values: dict[str, float]) -> list[str]:
    """Раскладывает суммы по строкам блока.

    Если фамилии в блоке нет, занимает первую свободную строку и вписывает
    её туда: блоки администраторов и кассиров пустые по замыслу, а состав
    официантов меняется. Возвращает тех, кому не хватило места.
    """
    start, end = block
    rows = {ws.cell(r, 1).value: r for r in range(start, end + 1)}
    no_room = []

    for name, amount in values.items():
        row = rows.get(name)
        if row is None:
            if not amount:
                continue  # нулю в отчёте делать нечего
            free = next(
                (r for r in range(start, end + 1) if not ws.cell(r, 1).value), None
            )
            if free is None:
                no_room.append(name)
                continue
            ws.cell(free, 1, name)
            rows[name] = row = free
        ws.cell(row, col, amount if amount else None)

    return no_room


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("export", help="файл выгрузки из приложения")
    ap.add_argument("--report", help="файл отчёта (по умолчанию ищется по месяцу)")
    args = ap.parse_args()

    data = json.loads(Path(args.export).read_text("utf-8"))
    day = data["businessDay"]

    if args.report:
        report = Path(args.report).expanduser()
        if not report.is_absolute():
            report = report_dir() / report
        if not report.exists():
            sys.exit(f"Не нашёл отчёт: {report}")
    else:
        months = "январь февраль март апрель май июнь июль август сентябрь октябрь ноябрь декабрь".split()
        year, month, _ = day.split("-")
        name = f"Чаевые-дубль-{months[int(month) - 1]}-{year}.xlsx"
        # Сначала там, куда кладём сегодня, потом в корне проекта: отчёты
        # за июль и август лежат в корне, они старше переезда на Диск.
        found = [p for p in (report_dir() / name, ROOT / name) if p.exists()]
        if not found:
            sys.exit(
                f"Не нашёл отчёт «{name}».\n"
                f"  Искал в: {report_dir()}\n"
                f"           {ROOT}\n"
                f"  Создать: python3 scripts/report_new.py {day[:7]}"
            )
        report = found[0]

    wb = load_workbook(report)
    ws = wb[wb.sheetnames[0]]
    col = find_date_column(ws, day)
    blocks = find_blocks(ws)

    missing: list[str] = []
    for title, field in WAITER_BLOCKS.items():
        if title not in blocks:
            sys.exit(
                f"В отчёте нет блока «{title}».\n"
                "  Отчёт собран старой версией шаблона. Создайте новый:\n"
                f"  python3 scripts/report_new.py {day[:7]}"
            )
        # Выгрузки, сделанные до появления вида денег, поля не содержат —
        # для них это просто ноль, а не повод останавливаться.
        missing += put(
            ws,
            blocks[title],
            col,
            {w["name"]: value_of(w, field) for w in data["waiters"]},
        )

    for title, who in ((ADMIN_BLOCK, "admin"), (CASHIER_BLOCK, "cashier")):
        person = data[who]
        if person["name"]:
            missing += put(ws, blocks[title], col, {person["name"]: person["amount"]})
        elif person["amount"]:
            print(
                f"  ⚠ {who}: {person['amount']:,.2f} ₽ не разнесены — "
                f"в смене не указан человек".replace(",", " ")
            )

    wb.save(report)

    print(f"Смена {day} записана в столбец {ws.cell(1, col).value}")
    print(f"Отчёт: {report.name}")
    print(f"Папка: {describe(report.parent)}")
    print(f"Официантов: {len(data['waiters'])}")
    t = data["totals"]
    # Виды денег перечисляем все до одного: строка «заработано = A + B + C»,
    # где сумма не сходится, читается как ошибка расчёта, хотя на деле
    # в ней просто забыли слагаемое.
    print(
        f"Заработано {t['earned']:,.2f} ₽ = наличные {t['cash']:,.2f} + "
        f"терминал {t['terminal']:,.2f} + QR {t.get('qr', 0):,.2f} + "
        f"сбор {t['service']:,.2f}".replace(",", " ")
    )
    diff = round(t["earned"] - t["payout"], 2)
    print(f"Распределено {t['payout']:,.2f} ₽, расхождение {diff:,.2f} ₽".replace(",", " "))
    if diff:
        print("  ⚠ расхождение не ноль — проверьте смену в приложении")

    if missing:
        print(f"\n⚠ Не нашёл в отчёте строки для: {', '.join(sorted(set(missing)))}")
        print("  Допишите фамилии в пустые строки нужных блоков и запустите снова.")
    print("\nФормулы пересчитаются при открытии файла.")


if __name__ == "__main__":
    main()
